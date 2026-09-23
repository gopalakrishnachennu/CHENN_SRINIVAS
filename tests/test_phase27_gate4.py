"""
Phase 2.7 Gate 4 — live harness, DOCX branding, UI auth, SQLite-only reads.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

from resume_engine.export.document_model import ContactHeader
from resume_engine.export.service import export_resume
from resume_engine.learning.repository import LearningRepository, reset_default_repository_for_tests
from resume_engine.learning.strategy_memory import load_outcome_records
from resume_engine.live.phase1_harness import preflight_live_harness
from resume_engine.models.resume_schema import ResumeExperience, ResumeJSON
from resume_engine.ui.app import create_app
from resume_engine.ui.auth import auth_enabled, verify_credentials


def _sample_resume() -> ResumeJSON:
    return ResumeJSON(
        target_title="Senior DevOps Engineer",
        summary="Platform engineer focused on AWS, Terraform, Kubernetes, and CI/CD delivery.",
        technical_skills={"Cloud": ["AWS", "Terraform"], "Platform": ["Kubernetes"]},
        experience=[
            ResumeExperience(
                company="Acme",
                title="DevOps Engineer",
                bullets=["Built CI/CD pipelines for Kubernetes workloads on AWS."],
            )
        ],
        projects=[],
        certifications=["AWS Solutions Architect"],
        variant_id="V01",
    )


# ---------------------------------------------------------------------------
# Live harness preflight (never spends money)
# ---------------------------------------------------------------------------


def test_live_preflight_skips_without_flag(monkeypatch):
    monkeypatch.delenv("RUN_LIVE_TESTS", raising=False)
    ok, reason = preflight_live_harness()
    assert ok is False
    assert "LIVE_TEST_NOT_RUN" in reason


def test_live_preflight_skips_without_api_key(monkeypatch):
    monkeypatch.setenv("RUN_LIVE_TESTS", "1")
    monkeypatch.setenv("OPENAI_API_KEY", "PASTE_YOUR_OPENAI_KEY_HERE")
    ok, reason = preflight_live_harness()
    assert ok is False
    assert "OPENAI_API_KEY" in reason


# ---------------------------------------------------------------------------
# SQLite-only default reads
# ---------------------------------------------------------------------------


def test_default_load_outcome_records_is_sqlite_only(tmp_path):
    jsonl = tmp_path / "orphan.jsonl"
    jsonl.write_text(
        '{"run_id":"orphan","variant_id":"V99","passed":true,"primary_family":"x",'
        '"secondary_family":"none","variant_positioning":"y","score_after":90}\n',
        encoding="utf-8",
    )
    repo = LearningRepository(db_path=tmp_path / "sot.sqlite3", dual_write_jsonl=False)
    reset_default_repository_for_tests(repo)
    # SQLite empty → default load returns [] (no JSONL fallback).
    assert load_outcome_records() == []
    repo.save_outcome(
        {
            "run_id": "sqlite-1",
            "jd_hash": "h",
            "variant_id": "V01",
            "primary_family": "devops_cloud",
            "secondary_family": "none",
            "seniority": "senior",
            "passed": True,
            "eligible_for_learning": True,
            "score_after": 95.0,
        }
    )
    records = load_outcome_records()
    assert any(r.get("run_id") == "sqlite-1" for r in records)
    # Explicit path still reads JSONL for tests.
    explicit = load_outcome_records(jsonl)
    assert any(r.get("run_id") == "orphan" for r in explicit)
    reset_default_repository_for_tests(None)


# ---------------------------------------------------------------------------
# DOCX branding
# ---------------------------------------------------------------------------


def test_docx_branding_uses_accent_heading_and_border(tmp_path: Path):
    export_resume(
        _sample_resume(),
        formats=("docx",),
        output_dir=tmp_path,
        basename="branded",
        contact=ContactHeader(name="Ada Lovelace", email="ada@example.com", location="London"),
    )
    doc = Document(str(tmp_path / "branded.docx"))
    texts = [p.text for p in doc.paragraphs]
    assert any("Ada Lovelace" in t for t in texts)
    assert any("PROFESSIONAL SUMMARY" in t.upper() for t in texts)
    # Accent color applied on a heading run.
    colored = False
    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.color and run.font.color.rgb and "SUMMARY" in (run.text or "").upper():
                colored = True
    assert colored


# ---------------------------------------------------------------------------
# UI auth
# ---------------------------------------------------------------------------


def test_auth_disabled_by_default(monkeypatch):
    monkeypatch.delenv("RESUME_ENGINE_UI_PASSWORD", raising=False)
    assert auth_enabled() is False
    client = create_app().test_client()
    response = client.get("/")
    assert response.status_code == 200
    assert b"Auth disabled" in response.data


def test_auth_redirects_when_password_configured(monkeypatch):
    monkeypatch.setenv("RESUME_ENGINE_UI_PASSWORD", "s3cret")
    monkeypatch.setenv("RESUME_ENGINE_UI_USER", "operator")
    client = create_app().test_client()
    response = client.get("/", follow_redirects=False)
    assert response.status_code in {302, 303}
    assert "/login" in (response.headers.get("Location") or "")


def test_auth_login_success(monkeypatch):
    monkeypatch.setenv("RESUME_ENGINE_UI_PASSWORD", "s3cret")
    monkeypatch.setenv("RESUME_ENGINE_UI_USER", "operator")
    assert verify_credentials("operator", "s3cret") is True
    assert verify_credentials("operator", "wrong") is False
    client = create_app().test_client()
    login_page = client.get("/login")
    assert login_page.status_code == 200
    html = login_page.data.decode("utf-8", errors="ignore")
    import re

    match = re.search(r'name="csrf_token"[^>]*value="([^"]+)"', html)
    assert match, "csrf token missing on login page"
    response = client.post(
        "/login",
        data={
            "username": "operator",
            "password": "s3cret",
            "next": "/",
            "csrf_token": match.group(1),
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Dashboard" in response.data or b"Control Center" in response.data
