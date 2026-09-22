"""DOCX exporter for validated ResumeJSON (Phase 3 / Gate 4 branding)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from resume_engine.export.document_model import (
    ContactHeader,
    ResumeDocumentView,
    build_document_view,
)
from resume_engine.models.resume_schema import ResumeJSON

# Brand accent aligned with operator UI teal (not purple/glow defaults).
ACCENT = RGBColor(0x0F, 0x5C, 0x4C)
INK = RGBColor(0x1C, 0x1A, 0x16)
MUTED = RGBColor(0x5C, 0x56, 0x4B)


def _set_run_font(
    run,
    *,
    bold: bool = False,
    size: int = 11,
    color: RGBColor | None = None,
    name: str = "Calibri",
) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color


def _set_paragraph_bottom_border(paragraph, color_hex: str = "0F5C4C", size: str = "12") -> None:
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def export_resume_docx(
    resume: ResumeJSON | dict,
    output_path: str | Path,
    *,
    contact: ContactHeader | dict | None = None,
) -> Path:
    view = build_document_view(resume, contact=contact)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    for section in document.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    name = view.contact.name or view.title
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run(name)
    _set_run_font(run, bold=True, size=18, color=INK, name="Calibri")

    if view.contact.name and view.title:
        sub = document.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.space_before = Pt(0)
        sub.paragraph_format.space_after = Pt(2)
        run = sub.add_run(view.title)
        _set_run_font(run, bold=False, size=12, color=ACCENT)

    contact_line = view.contact.contact_line()
    brand_anchor = document.add_paragraph()
    brand_anchor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand_anchor.paragraph_format.space_before = Pt(0)
    brand_anchor.paragraph_format.space_after = Pt(8)
    if contact_line:
        run = brand_anchor.add_run(contact_line)
        _set_run_font(run, size=9, color=MUTED)
    else:
        run = brand_anchor.add_run(" ")
        _set_run_font(run, size=9)
    _set_paragraph_bottom_border(brand_anchor, color_hex="0F5C4C", size="18")

    def add_heading(text: str) -> None:
        p = document.add_paragraph()
        run = p.add_run(text.upper())
        _set_run_font(run, bold=True, size=11, color=ACCENT)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        _set_paragraph_bottom_border(p, color_hex="0F5C4C", size="6")

    if view.summary:
        add_heading("Professional Summary")
        p = document.add_paragraph(view.summary)
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            _set_run_font(run, size=10, color=INK)

    if view.technical_skills:
        add_heading("Technical Skills")
        for category, skills in view.technical_skills.items():
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            label = p.add_run(f"{category}: ")
            _set_run_font(label, bold=True, size=10, color=INK)
            values = p.add_run(", ".join(skills))
            _set_run_font(values, size=10, color=INK)

    if view.experience:
        add_heading("Experience")
        for job in view.experience:
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(f"{job['title']} — {job['company']}")
            _set_run_font(run, bold=True, size=10, color=INK)
            for bullet in job.get("bullets") or []:
                bp = document.add_paragraph(bullet, style="List Bullet")
                bp.paragraph_format.space_after = Pt(1)
                for run in bp.runs:
                    _set_run_font(run, size=10, color=INK)

    if view.projects:
        add_heading("Projects")
        for project in view.projects:
            p = document.add_paragraph()
            run = p.add_run(project["name"])
            _set_run_font(run, bold=True, size=10, color=INK)
            if project.get("summary"):
                sp = document.add_paragraph(project["summary"])
                for run in sp.runs:
                    _set_run_font(run, size=10, color=INK)
            tech = project.get("technologies") or []
            if tech:
                tp = document.add_paragraph()
                label = tp.add_run("Technologies: ")
                _set_run_font(label, bold=True, size=10, color=ACCENT)
                values = tp.add_run(", ".join(tech))
                _set_run_font(values, size=10, color=INK)
            for bullet in project.get("bullets") or []:
                bp = document.add_paragraph(bullet, style="List Bullet")
                for run in bp.runs:
                    _set_run_font(run, size=10, color=INK)

    if view.certifications:
        add_heading("Certifications")
        for cert in view.certifications:
            bp = document.add_paragraph(cert, style="List Bullet")
            for run in bp.runs:
                _set_run_font(run, size=10, color=INK)

    document.save(str(path))
    return path


__all__ = ["export_resume_docx", "ResumeDocumentView"]
