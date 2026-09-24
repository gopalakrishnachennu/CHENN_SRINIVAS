#!/usr/bin/env python3
"""Import resume_engine_candidate_jd_data.xlsx into SQLite UI tables.

Replaces dummy/test candidate profiles and JD library rows with workbook data.
Does not call OpenAI — jobs are stored with JD text ready for Analyze later.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from resume_engine.config.settings import PROJECT_ROOT
from resume_engine.ui.db import connect_ui_db, ensure_ui_schema
from resume_engine.ui.services.family_registry_service import resolve_family

DEFAULT_XLSX = PROJECT_ROOT / "data" / "resume_engine_candidate_jd_data.xlsx"


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _fmt_date(value) -> str:
    if value is None:
        return "Present"
    if hasattr(value, "strftime"):
        return value.strftime("%b %Y")
    text = str(value).strip()
    return text or "Present"


def _sheet_dicts(wb, name: str) -> list[dict]:
    ws = wb[name]
    headers = [c.value for c in ws[1]]
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw[0] is None:
            continue
        rows.append({headers[i]: raw[i] for i in range(len(headers)) if headers[i]})
    return rows


def clear_dummy_data(conn) -> dict[str, int]:
    """Remove test/dummy operator data; keep schema and family registry."""
    counts = {}
    for table in (
        "match_cache",
        "resume_library",
        "resume_runs",
        "jd_library",
        "candidate_profile_versions",
        "candidate_profiles",
    ):
        cur = conn.execute(f"SELECT COUNT(*) AS n FROM {table}")
        counts[table] = int(cur.fetchone()["n"])
        conn.execute(f"DELETE FROM {table}")
    return counts


def import_workbook(xlsx: Path, *, clear: bool = True) -> dict:
    ensure_ui_schema()
    wb = load_workbook(xlsx, data_only=True)
    candidates = _sheet_dicts(wb, "Candidates")
    employment = _sheet_dicts(wb, "Employment History")
    jobs = _sheet_dicts(wb, "Jobs")
    matches = _sheet_dicts(wb, "Matches")

    emp_by_id: dict[str, list] = {}
    for row in employment:
        cid = str(row["Candidate_ID"])
        emp_by_id.setdefault(cid, []).append(row)

    now = _now()
    with connect_ui_db() as conn:
        cleared = clear_dummy_data(conn) if clear else {}

        # --- Candidates ---
        for row in candidates:
            cid = str(row["Candidate_ID"])
            name = str(row["Candidate_Name"]).strip()
            primary = resolve_family(row.get("Primary_Family")) or ""
            secondary = resolve_family(row.get("Secondary_Family")) or ""
            cert = row.get("Certificates_Optional")
            companies = []
            for emp in emp_by_id.get(cid, []):
                end = emp.get("End_Date")
                companies.append(
                    {
                        "company": str(emp.get("Company") or "").strip(),
                        "start_date": _fmt_date(emp.get("Start_Date")),
                        "end_date": "Present" if end in (None, "") else _fmt_date(end),
                        "provenance": {
                            "company": "VERIFIED_CANDIDATE_INPUT",
                            "start_date": "VERIFIED_CANDIDATE_INPUT",
                            "end_date": "VERIFIED_CANDIDATE_INPUT",
                            "role_title": "NOT_PROVIDED",
                        },
                    }
                )
            payload = {
                "candidate_name": name,
                "email": "",
                "phone": "",
                "location": str(row.get("Location") or ""),
                "linkedin": "",
                "website": "",
                "primary_family": primary,
                "secondary_family": secondary,
                "companies": companies,
                "education": [],
                "certifications": [str(cert).strip()] if cert else [],
                "_schema": "candidate_facts_v2",
                "_external_id": cid,
            }
            conn.execute(
                """
                INSERT INTO candidate_profiles(id, name, status, payload_json, created_at, updated_at)
                VALUES (?, ?, 'active', ?, ?, ?)
                """,
                (cid, name, json.dumps(payload), now, now),
            )
            conn.execute(
                """
                INSERT INTO candidate_profile_versions(profile_id, version, payload_json, actor, created_at)
                VALUES (?, 1, ?, 'workbook_import', ?)
                """,
                (cid, json.dumps(payload), now),
            )

        # --- Jobs ---
        for row in jobs:
            jid = str(row["JD_ID"])
            primary = resolve_family(row.get("Primary_Family")) or ""
            secondary = resolve_family(row.get("Secondary_Family")) or ""
            if secondary in {"none", "null"}:
                secondary = None
            meta = {
                "external_id": jid,
                "key_responsibilities": row.get("Key_Responsibilities"),
                "must_have": row.get("Must_Have"),
                "preferred": row.get("Preferred"),
                "posting_status": row.get("Posting_Status"),
                "verified_on": str(row.get("Verified_On") or ""),
                "from_workbook": True,
            }
            # Split must-have / preferred into lists for UI display helpers
            must = [s.strip() for s in str(row.get("Must_Have") or "").split(";") if s.strip()]
            pref = [s.strip() for s in str(row.get("Preferred") or "").split(";") if s.strip()]
            meta["p1"] = must[:3]
            meta["p2"] = must[3:]
            meta["p3"] = pref
            meta["display"] = {
                "must_have": must,
                "preferred": pref,
                "responsibilities": [
                    s.strip()
                    for s in str(row.get("Key_Responsibilities") or "").split(";")
                    if s.strip()
                ],
            }
            conn.execute(
                """
                INSERT INTO jd_library(
                  id, jd_hash, title, company, location, job_url, source, seniority,
                  primary_family, secondary_family, status, blueprint_path, jd_text,
                  metadata_json, created_at, updated_at,
                  analysis_status, analysis_version, jd_content_hash
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?, ?, ?, ?, 'NEEDS_ANALYSIS', 0, ?)
                """,
                (
                    jid,
                    jid,
                    str(row.get("Job_Title") or "Untitled"),
                    str(row.get("Company") or ""),
                    str(row.get("Location") or ""),
                    str(row.get("Source_URL") or ""),
                    "workbook",
                    str(row.get("Seniority") or ""),
                    primary,
                    secondary,
                    None,  # no blueprint until analyzed
                    str(row.get("JD_Input_Text") or ""),
                    json.dumps(meta),
                    now,
                    now,
                    None,
                ),
            )

        # --- Match cache (workbook labels → engine types) ---
        label_map = {
            "DIRECT_MATCH": "DIRECT",
            "HYBRID_MATCH": "HYBRID",
            "SECONDARY_MATCH": "SECONDARY",
            "COMPATIBLE_MATCH": "COMPATIBLE",
            "NO_MATCH": "NO_MATCH",
        }
        score_map = {"DIRECT": 1.0, "HYBRID": 0.85, "SECONDARY": 0.7, "COMPATIBLE": 0.5, "NO_MATCH": 0.0}
        for row in matches:
            mt = label_map.get(str(row.get("Match_Type") or ""), "NO_MATCH")
            if mt == "NO_MATCH":
                continue  # product UI shows matches only
            cid = str(row["Candidate_ID"])
            jid = str(row["JD_ID"])
            details = {
                "match_type": mt,
                "score": score_map[mt],
                "reason": "workbook_import",
                "candidate_id": cid,
                "job_id": jid,
                "job_title": row.get("Job_Title"),
                "job_company": row.get("Job_Company"),
                "create_resume": row.get("Create_Resume"),
                "priority": row.get("Priority"),
            }
            conn.execute(
                """
                INSERT INTO match_cache(candidate_id, job_id, match_type, score, details_json, created_at)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(candidate_id, job_id) DO UPDATE SET
                  match_type=excluded.match_type,
                  score=excluded.score,
                  details_json=excluded.details_json,
                  created_at=excluded.created_at
                """,
                (cid, jid, mt, score_map[mt], json.dumps(details), now),
            )

        conn.commit()

    return {
        "xlsx": str(xlsx),
        "cleared": cleared,
        "candidates": len(candidates),
        "employment_rows": len(employment),
        "jobs": len(jobs),
        "matches_stored": sum(
            1
            for m in matches
            if label_map.get(str(m.get("Match_Type") or ""), "NO_MATCH") != "NO_MATCH"
        ),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=DEFAULT_XLSX,
        help="Path to workbook (default: data/resume_engine_candidate_jd_data.xlsx)",
    )
    parser.add_argument(
        "--keep-existing",
        action="store_true",
        help="Do not clear existing candidates/jobs before import",
    )
    args = parser.parse_args(argv)
    if not args.xlsx.exists():
        print(f"Workbook not found: {args.xlsx}", file=sys.stderr)
        return 2
    result = import_workbook(args.xlsx, clear=not args.keep_existing)
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
